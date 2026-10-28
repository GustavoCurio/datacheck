"""Leitura limitada de CSV e XLSX sem executar fórmulas."""
import csv
import io
from numbers import Number
from datetime import date, datetime
from pathlib import Path
from zipfile import ZipFile

import pandas as pd

MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 50_000
MAX_COLS = 50
MAX_CELLS = 500_000
MAX_CELL_CHARS = 2_000


def _frame(rows):
    if not rows or not rows[0]:
        raise ValueError("O arquivo está vazio ou não tem cabeçalho.")
    headers = [str(x).strip() if x is not None else "" for x in rows[0]]
    if len(headers) > MAX_COLS or any(not h or len(h) > 100 for h in headers):
        raise ValueError("Use até 50 colunas, com nomes preenchidos de até 100 caracteres.")
    if len(set(headers)) != len(headers):
        raise ValueError("Os nomes das colunas precisam ser únicos.")
    if not rows[1:]:
        raise ValueError("O arquivo precisa ter pelo menos um registro além do cabeçalho.")
    if len(rows) - 1 > MAX_ROWS or (len(rows) - 1) * len(headers) > MAX_CELLS:
        raise ValueError("Limite: 50.000 registros e 500.000 células por arquivo.")
    converted = []
    native_numeric = []
    for position, row in enumerate(rows[1:]):
        if len(row) != len(headers):
            raise ValueError("Existe um registro com quantidade de campos diferente do cabeçalho.")
        values = []
        for column, value in zip(headers, row):
            if isinstance(value, Number) and not isinstance(value, bool):
                native_numeric.append((position, column))
            if isinstance(value, (datetime, date)):
                value = value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
            value = "" if value is None else str(value)
            if len(value) > MAX_CELL_CHARS:
                raise ValueError("Uma célula ultrapassa o limite de 2.000 caracteres.")
            values.append(value)
        converted.append(values)
    table = pd.DataFrame(converted, columns=headers, dtype=str)
    table.attrs["native_numeric"] = native_numeric
    return table


def read_table(content: bytes, filename: str, delimiter="auto", encoding="utf-8-sig", sheet=0):
    if not content or len(content) > MAX_BYTES:
        raise ValueError("Envie um arquivo de até 10 MB com conteúdo.")
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        if delimiter not in {"auto", ",", ";", "\t"} or encoding not in {"utf-8-sig", "cp1252"}:
            raise ValueError("Separador ou codificação inválidos.")
        try:
            text = content.decode(encoding)
            if "\x00" in text:
                raise ValueError("CSV contém caracteres nulos.")
            if delimiter == "auto":
                try:
                    delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",;\t").delimiter
                except csv.Error:
                    delimiter = ","
            reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True)
            rows = []
            for row in reader:
                if not row:
                    continue
                rows.append(row)
                if len(rows) > MAX_ROWS + 1:
                    raise ValueError("Limite de 50.000 registros por arquivo.")
            return _frame(rows)
        except (UnicodeDecodeError, csv.Error) as error:
            raise ValueError("Não foi possível ler o CSV. Confira a codificação e o separador.") from error
    if suffix != ".xlsx":
        raise ValueError("Formatos aceitos: CSV e XLSX.")
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            if len(archive.infolist()) > 1000 or sum(item.file_size for item in archive.infolist()) > 30 * 1024 * 1024:
                raise ValueError("A planilha excede o limite descompactado de 30 MB ou 1.000 entradas.")
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
        try:
            if type(sheet) is not int or sheet < 0 or sheet >= len(workbook.worksheets):
                raise ValueError("A aba selecionada não existe.")
            selected = workbook.worksheets[sheet]
            if (selected.max_row or 0) > MAX_ROWS + 1 or (selected.max_column or 0) > MAX_COLS:
                raise ValueError("A aba excede 50.000 registros ou 50 colunas, incluindo células formatadas.")
            if max((selected.max_row or 0) - 1, 0) * (selected.max_column or 0) > MAX_CELLS:
                raise ValueError("Limite de 500.000 células por aba.")
            rows = []
            for cells in selected.iter_rows():
                if any(cell.data_type == "f" for cell in cells):
                    raise ValueError("A aba contém fórmulas. Exporte os valores para CSV ou XLSX sem fórmulas.")
                rows.append([cell.value for cell in cells])
            return _frame(rows)
        finally:
            workbook.close()
    except ValueError:
        raise
    except Exception as error:
        raise ValueError("Não foi possível ler o XLSX. Confira se o arquivo é uma planilha válida.") from error
